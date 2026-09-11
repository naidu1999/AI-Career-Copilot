from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

def to_docx(title:str,content:str)->bytes:
    document=Document();document.add_heading(title,0)
    for line in content.splitlines():
        line=line.strip()
        if not line:continue
        if line.startswith("#"):document.add_heading(line.lstrip("# ").strip(),level=min(3,len(line)-len(line.lstrip("#"))))
        elif line.startswith(("- ","* ")):document.add_paragraph(line[2:],style="List Bullet")
        else:document.add_paragraph(line)
    stream=BytesIO();document.save(stream);return stream.getvalue()

def to_pdf(title:str,content:str)->bytes:
    stream=BytesIO();styles=getSampleStyleSheet();story=[Paragraph(escape(title),styles["Title"]),Spacer(1,12)]
    for line in content.splitlines():
        if line.strip():story.extend([Paragraph(escape(line.strip()),styles["BodyText"]),Spacer(1,6)])
    SimpleDocTemplate(stream,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story)
    return stream.getvalue()

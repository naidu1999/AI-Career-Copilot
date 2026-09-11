import hashlib
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def valid_signature(data:bytes,content_type:str)->bool:
    if content_type=="application/pdf":return data.startswith(b"%PDF-")
    if content_type=="application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return data.startswith(b"PK\x03\x04")
    return False


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == "application/pdf":
        reader = PdfReader(BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(BytesIO(data))
        blocks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            blocks.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        return "\n".join(blocks).strip()
    raise ValueError("Unsupported document type")


def safe_filename(name: str) -> str:
    return Path(name or "resume").name.replace(" ", "_")

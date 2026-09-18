from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app
from backend.jobs.matcher import match_job
from backend.jobs.providers import canonical_key

def profile(**extra):
    base={"target_titles":["Data Analyst"],"skills":["Python","SQL","Power BI"],"locations":["Bengaluru"],
          "remote_allowed":True,"remote_countries":["India"],"years_experience":2,"relevant_experience":2,
          "excluded_roles":[],"excluded_employment_types":[],"verified_evidence_ratio":1}
    return base|extra

def test_unrelated_roles_cannot_score_high():
    result=match_job(profile(),{"title":"Senior Sales Manager","description":"CRM, sales targets, 8+ years","location":"Bengaluru","employment_type":"full-time"})
    assert result["classification"]=="ineligible"
    assert result["score"]<=35
    assert any("Unrelated role family" in x for x in result["blockers"])

def test_related_role_is_reviewable_but_not_full_title_match():
    result=match_job(profile(),{"title":"Data Engineer","description":"Python SQL ETL, 2 years","location":"Bengaluru","employment_type":"full-time"})
    assert result["components"]["title"]==55
    assert result["score"]<90

def test_canonical_duplicate_identity_ignores_provider_description():
    one=canonical_key("Example Pvt Ltd","Data Analyst II","Bengaluru Hybrid")
    two=canonical_key("Example","Data Analyst 2","Bengaluru")
    assert one!=canonical_key("Other","Data Analyst","Bengaluru")
    assert len(one)==64 and len(two)==64

@pytest.mark.asyncio
async def test_scan_merges_same_vacancy_across_sources(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    from backend.db.local import execute,initialize,now,rows,uid
    from backend.jobs.providers import Job,key
    import backend.services.discovery as discovery
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"dedupe.db"));initialize()
    for name,provider in (("Official","greenhouse"),("Aggregator","jooble")):
        execute("INSERT INTO job_sources (id,name,provider,board_key,enabled,created_at,owner_id) VALUES (?,?,?,?,?,?,?)",(uid(),name,provider,name.lower(),1,now(),"default"))
    async def fake(provider,board,company,p):
        return [Job(key(provider,board),provider,board,"Data Analyst","Example Ltd","Bengaluru","Python SQL","https://example.test/"+provider,now(),"posted")]
    monkeypatch.setattr(discovery,"fetch_jobs",fake)
    result=await discovery.scan(profile(id="default"),"test")
    assert result["new_jobs"]==1 and result["duplicates"]==1
    assert rows("SELECT COUNT(*) n FROM jobs")[0]["n"]==1
    assert rows("SELECT COUNT(*) n FROM job_source_refs")[0]["n"]==2

def docx_bytes()->bytes:
    doc=Document();doc.add_paragraph("Rahul Naidu");doc.add_heading("Skills",1);doc.add_paragraph("Python, SQL, Power BI")
    stream=BytesIO();doc.save(stream);return stream.getvalue()

def test_end_to_end_upload_evidence_match_save_status(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"e2e.db"));monkeypatch.setattr(settings,"UPLOAD_DIR",str(tmp_path/"uploads"))
    monkeypatch.setattr(settings,"BACKUP_DIR",str(tmp_path/"backups"));monkeypatch.setattr(settings,"SCAN_ENABLED",False)
    with TestClient(app) as client:
        upload=client.post("/api/resumes",files={"file":("resume.docx",docx_bytes(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert upload.status_code==201
        evidence=client.get("/api/evidence").json();assert evidence
        assert client.patch(f"/api/evidence/{evidence[0]['id']}",json={"status":"verified"}).status_code==200
        job={"title":"Data Analyst","company":"Example","location":"Bengaluru","description":"Python SQL Power BI, 2 years","url":"https://example.test/jobs/analyst","employment_type":"full-time"}
        assert client.post("/api/jobs/manual",json=job).status_code==201
        from time import monotonic, sleep
        matches=[]
        deadline=monotonic()+10
        while monotonic()<deadline:
            matches=client.get("/api/jobs?classification=all&strict_date=false&hours=720").json()["items"]
            if matches:break
            sleep(0.2)
        assert matches,"background re-match should produce the manual job's match"
        saved=client.post("/api/applications",json={"job_id":matches[0]["id"],"status":"shortlisted"}).json()
        update={"status":"interview","notes":"Round one","recruiter_name":"A","recruiter_contact":"","follow_up_at":None,"interview_at":None,"deadline_at":None,"rejection_reason":"","salary_details":"","offer_details":""}
        assert client.patch(f"/api/applications/{saved['id']}",json=update).status_code==200
        history=client.get(f"/api/applications/{saved['id']}/history").json()
        assert history[0]["to_status"]=="interview" and len(history)>=2

def test_upload_signature_mismatch_is_rejected(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"signature.db"));monkeypatch.setattr(settings,"UPLOAD_DIR",str(tmp_path/"uploads"))
    with TestClient(app) as client:
        response=client.post("/api/resumes",files={"file":("fake.pdf",b"not a pdf","application/pdf")})
        assert response.status_code==415

def test_verified_backup_restore_round_trip(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"restore.db"));monkeypatch.setattr(settings,"UPLOAD_DIR",str(tmp_path/"uploads"));monkeypatch.setattr(settings,"BACKUP_DIR",str(tmp_path/"backups"))
    with TestClient(app) as client:
        p=client.get("/api/profile").json();p["full_name"]="Before";client.put("/api/profile",json=p)
        created=client.post("/api/backup").json();assert created["valid"]
        p["full_name"]="After";client.put("/api/profile",json=p)
        restored=client.post("/api/backups/restore",json={"path":created["path"],"confirmation":"RESTORE BACKUP"})
        assert restored.status_code==200
        assert client.get("/api/profile").json()["full_name"]=="Before"

@pytest.mark.asyncio
async def test_ai_router_falls_back_and_records_health(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    from backend.db.local import initialize,rows
    from backend.services.ai_router import AIRouter,Provider
    import backend.services.ai_router as module
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"router.db"))
    providers=[Provider("one","One","openai","http://one","x","m1"),Provider("two","Two","openai","http://two","x","m2")]
    monkeypatch.setattr(module,"configured_providers",lambda:providers)
    async def fake(provider,messages):
        if provider.name=="One":raise __import__("httpx").TimeoutException("timeout")
        return "valid response"
    monkeypatch.setattr(module,"request_provider",fake);initialize()
    result=await AIRouter("default").chat("summary",[{"role":"user","content":"hello"}])
    assert result["provider"]=="Two" and result["attempts"]==2
    assert rows("SELECT COUNT(*) n FROM ai_request_log WHERE status='success'")[0]["n"]==1

@pytest.mark.asyncio
async def test_keyless_pollinations_provider_is_used_without_api_keys(tmp_path,monkeypatch):
    from backend.core.config_new import settings
    from backend.db.local import initialize,rows
    import backend.services.ai_router as module
    from backend.services.ai_router import AIRouter
    monkeypatch.setattr(settings,"DATABASE_PATH",str(tmp_path/"keyless.db"))
    for field in ("AI_API_KEY","OPENAI_API_KEY","ANTHROPIC_API_KEY","GROQ_API_KEY","GEMINI_API_KEY","OPENROUTER_API_KEY","CEREBRAS_API_KEY","MISTRAL_API_KEY","AI_MODEL","OLLAMA_MODEL"):
        monkeypatch.setattr(settings,field,"")
    monkeypatch.setattr(settings,"AI_PROVIDER","rules");monkeypatch.setattr(settings,"POLLINATIONS_ENABLED",True)
    async def fake(provider,messages):
        assert provider.key=="" and provider.base_url.endswith("pollinations.ai/openai")
        return "keyless response"
    monkeypatch.setattr(module,"request_provider",fake);initialize()
    result=await AIRouter("default").chat("summary",[{"role":"user","content":"hello"}])
    assert result["provider"].startswith("Pollinations")
    assert rows("SELECT COUNT(*) n FROM ai_request_log WHERE status='success'")[0]["n"]==1

def test_accessibility_and_motion_contract():
    html=open("backend/static/index.html",encoding="utf-8").read();css=open("backend/static/style.css",encoding="utf-8").read()
    assert 'aria-label="Primary"' in html and 'aria-live="polite"' in html
    assert "prefers-reduced-motion" in css and ".score-ring" in css and ".skeleton-card" in css
